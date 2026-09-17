from odoo import models, fields, api
from odoo.exceptions import ValidationError

class VendorComparison(models.Model):
    _name = "vendor.comparison"
    _description = "Vendor Comparison"
    _inherit = ['mail.thread']
    _rec_name = 'name'
    _order = 'create_date desc'

    name = fields.Char(string="Reference", copy=False)
    line_ids = fields.One2many('vendor.comparison.line', 'comparison_id', string="Products")
    state = fields.Selection([
        ('draft', 'Draft'),
        ('rfq_created', 'RFQ Created'),
    ], default='draft', tracking=True)
    reason = fields.Char()
    purchase_count = fields.Integer(string="Purchase Count",compute='_compute_purchase_order')
    vendor_count = fields.Integer(string="Vendor Count",compute='_compute_vendor_count',store=True)
    product_count = fields.Integer(string="Product Count",compute='_compute_product_count',store=True)
    vendor_comparison_table = fields.Html(string="Vendor Comparison Table", compute='_compute_get_purchase')

    @api.depends('line_ids.product_id', 'line_ids.partner_ids', 'state')
    def _compute_get_purchase(self):
        for rec in self:
            get_purchase = self.env['purchase.order.line'].search([
                ('order_id.vendor_comparison_id', '=', rec.id)
            ])
            product_dict = {}
            all_partners = []

            for pol in get_purchase:
                product = pol.product_id
                partner = pol.order_id.partner_id

                if product not in product_dict:
                    product_dict[product] = {
                        'uom': pol.product_uom_id.name or '',
                        'vendors': {}
                    }

                product_dict[product]['vendors'][partner] = {
                    'price': pol.price_unit,
                    'qty': pol.product_qty,
                }

                if partner not in all_partners:
                    all_partners.append(partner)

            # Color palette - light & medium
            HEADER_MAIN = "#4a90a4"  # medium teal
            HEADER_SUB = "#a8d5e2"  # light teal
            HEADER_TEXT = "#ffffff"
            SUBROW_TEXT = "#2c3e50"
            QTY_PRICE_COL = "#7b9ec4"  # medium blue
            QTY_PRICE_SUB = "#c5d9f1"  # light blue
            TOTAL_COL = "#e8a87c"  # medium peach
            TOTAL_SUB = "#fde8d8"  # light peach
            ROW_ODD = "#f4f9fb"  # very light teal
            ROW_EVEN = "#ffffff"
            FOOTER_BG = "#eaf4f8"
            PRICE_COLOR = "#2980b9"  # medium blue for price values
            QTY_COLOR = "#27ae60"  # medium green for qty values
            QXPRICE_COLOR = "#6c5ce7"  # medium purple for qty x price

            vendor_headers = ""
            for partner in all_partners:
                vendor_headers += f"""
                    <th colspan="3" style="padding:10px 16px; background-color:{HEADER_MAIN};
                        color:{HEADER_TEXT}; border:1px solid #b0cfe0; text-align:center;
                        white-space:nowrap; letter-spacing:0.5px;">{partner.name}</th>
                """

            sub_headers = f"""
                <th style="padding:8px 12px; background-color:{HEADER_SUB}; color:{SUBROW_TEXT}; border:1px solid #b0cfe0; text-align:center; font-weight:600;">Product</th>
                <th style="padding:8px 12px; background-color:{HEADER_SUB}; color:{SUBROW_TEXT}; border:1px solid #b0cfe0; text-align:center; font-weight:600;">UoM</th>
            """
            for _ in all_partners:
                sub_headers += f"""
                    <th style="padding:8px 12px; background-color:{HEADER_SUB}; color:{SUBROW_TEXT}; border:1px solid #b0cfe0; text-align:center; font-weight:600;">Price</th>
                    <th style="padding:8px 12px; background-color:{HEADER_SUB}; color:{SUBROW_TEXT}; border:1px solid #b0cfe0; text-align:center; font-weight:600;">Qty</th>
                    <th style="padding:8px 12px; background-color:{QTY_PRICE_SUB}; color:{SUBROW_TEXT}; border:1px solid #b0cfe0; text-align:center; font-weight:600;">Qty × Price</th>
                """
            sub_headers += f"""
                <th style="padding:8px 12px; background-color:{TOTAL_SUB}; color:{SUBROW_TEXT}; border:1px solid #f0c8a0; text-align:center; font-weight:600;">Grand Total</th>
            """

            vendor_totals = {partner: 0.0 for partner in all_partners}
            rows_html = ""
            grand_total = 0.0

            for i, (product, data) in enumerate(product_dict.items()):
                bg = ROW_ODD if i % 2 == 0 else ROW_EVEN
                vendors = data['vendors']
                row_grand_total = sum(v['qty'] * v['price'] for v in vendors.values())
                grand_total += row_grand_total

                row = f"""
                    <td style="padding:10px 14px; border:1px solid #d6eaf3; text-align:left;
                        font-weight:600; white-space:nowrap; color:{SUBROW_TEXT};">{product.name}</td>
                    <td style="padding:10px 14px; border:1px solid #d6eaf3; text-align:center;
                        color:#7f8c8d;">{data['uom']}</td>
                """

                for partner in all_partners:
                    if partner in vendors:
                        price = vendors[partner]['price']
                        qty = vendors[partner]['qty']
                        qty_x_price = qty * price
                        vendor_totals[partner] += qty_x_price
                        row += f"""
                            <td style="padding:10px; border:1px solid #d6eaf3; text-align:center;
                                color:{PRICE_COLOR}; font-weight:600;">{price:.2f}</td>
                            <td style="padding:10px; border:1px solid #d6eaf3; text-align:center;
                                color:{QTY_COLOR}; font-weight:600;">{int(qty)}</td>
                            <td style="padding:10px; border:1px solid #d6eaf3; text-align:center;
                                color:{QXPRICE_COLOR}; font-weight:600;">{qty_x_price:.2f}</td>
                        """
                    else:
                        row += """
                            <td style="padding:10px; border:1px solid #d6eaf3; text-align:center; color:#bdc3c7;">-</td>
                            <td style="padding:10px; border:1px solid #d6eaf3; text-align:center; color:#bdc3c7;">-</td>
                            <td style="padding:10px; border:1px solid #d6eaf3; text-align:center; color:#bdc3c7;">-</td>
                        """

                row += f"""
                    <td style="padding:10px; border:1px solid #f0c8a0; text-align:center;
                        color:#c0392b; font-weight:700; background-color:#fef5ec;">{row_grand_total:.2f}</td>
                """
                rows_html += f"<tr style='background-color:{bg};'>{row}</tr>"

            # Footer total row
            vendor_total_cols = ""
            for partner in all_partners:
                vendor_total_cols += f"""
                    <td style="border:1px solid #b0cfe0; background-color:{FOOTER_BG};"></td>
                    <td style="border:1px solid #b0cfe0; background-color:{FOOTER_BG};"></td>
                    <td style="padding:10px; border:1px solid #b0cfe0; text-align:center;
                        font-weight:700; color:{QXPRICE_COLOR}; background-color:{FOOTER_BG};">{vendor_totals[partner]:.2f}</td>
                """

            total_row = f"""
                <tr>
                    <td colspan="2" style="padding:10px 16px; border:1px solid #b0cfe0;
                        text-align:center; font-weight:700; color:#2c3e50;
                        background-color:{FOOTER_BG}; font-size:15px;">Total</td>
                    {vendor_total_cols}
                    <td style="padding:10px; border:1px solid #f0c8a0; text-align:center;
                        font-weight:700; color:#c0392b; background-color:#fde8d8;
                        font-size:15px;">{grand_total:.2f}</td>
                </tr>
            """

            full_message = f"""
                <div style="font-family:'Segoe UI', Arial, sans-serif; font-size:13px;
                    color:#2c3e50; overflow-x:auto; width:100%; padding:8px 0;">
                    <table style="min-width:900px; border-collapse:collapse;
                        border:1px solid #b0cfe0; border-radius:8px; overflow:hidden;">
                        <thead>
                            <tr>
                                <th colspan="2" style="padding:12px 16px; background-color:{HEADER_MAIN};
                                    color:{HEADER_TEXT}; border:1px solid #b0cfe0; text-align:center;
                                    white-space:nowrap; font-size:14px; letter-spacing:0.5px;">Product Info</th>
                                {vendor_headers}
                                <th style="padding:12px 16px; background-color:{TOTAL_COL};
                                    color:{HEADER_TEXT}; border:1px solid #f0c8a0; text-align:center;
                                    white-space:nowrap; font-size:14px; letter-spacing:0.5px;">Grand Total</th>
                            </tr>
                            <tr>{sub_headers}</tr>
                        </thead>
                        <tbody>
                            {rows_html}
                            {total_row}
                        </tbody>
                    </table>
                </div>
            """
            rec.vendor_comparison_table = full_message

    def action_export_vendor_comparison_xlsx(self):
        import io
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
        import base64

        for rec in self:
            get_purchase = self.env['purchase.order.line'].search([
                ('order_id.vendor_comparison_id', '=', rec.id)
            ])

            product_dict = {}
            all_partners = []

            for pol in get_purchase:
                product = pol.product_id
                partner = pol.order_id.partner_id

                if product not in product_dict:
                    product_dict[product] = {
                        'uom': pol.product_uom_id.name or '',
                        'vendors': {}
                    }

                product_dict[product]['vendors'][partner] = {
                    'price': pol.price_unit,
                    'qty': pol.product_qty,
                }

                if partner not in all_partners:
                    all_partners.append(partner)

            # ─── Colors ─────────────────────────────────────────────────────────
            C_HEADER_MAIN = "4A90A4"
            C_HEADER_SUB = "A8D5E2"
            C_QTY_PRICE_SUB = "C5D9F1"
            C_TOTAL_COL = "E8A87C"
            C_TOTAL_SUB = "FDE8D8"
            C_ROW_ODD = "F4F9FB"
            C_ROW_EVEN = "FFFFFF"
            C_FOOTER_BG = "EAF4F8"
            C_PRICE = "2980B9"
            C_QTY = "27AE60"
            C_QXPRICE = "6C5CE7"
            C_GRAND_TOTAL = "C0392B"
            C_WHITE = "FFFFFF"
            C_DARK = "2C3E50"
            C_GRAY = "7F8C8D"

            def mk_fill(hex_color):
                return PatternFill("solid", start_color=hex_color, fgColor=hex_color)

            def mk_font(bold=False, color="2C3E50", size=11):
                return Font(name="Segoe UI", bold=bold, color=color, size=size)

            def mk_align(h="center", v="center"):
                return Alignment(horizontal=h, vertical=v, wrap_text=False)

            def mk_border(color="B0CFE0"):
                s = Side(style="thin", color=color)
                return Border(left=s, right=s, top=s, bottom=s)

            def style(cell, bg=None, fnt=None, aln=None, brd=None):
                if bg:  cell.fill = mk_fill(bg)
                if fnt: cell.font = fnt
                if aln: cell.alignment = aln
                if brd: cell.border = brd

            # ─── Workbook ────────────────────────────────────────────────────────
            wb = Workbook()
            ws = wb.active
            ws.title = "Vendor Comparison"

            NUM_VENDOR_COLS = 3  # Price | Qty | Qty×Price
            start_col = 3
            total_col = start_col + (len(all_partners) * NUM_VENDOR_COLS)

            # ─── Row 1: Main Headers ─────────────────────────────────────────────
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2)
            c = ws.cell(row=1, column=1, value="Product Info")
            style(c, bg=C_HEADER_MAIN, fnt=mk_font(bold=True, color=C_WHITE, size=12),
                  aln=mk_align(), brd=mk_border())
            ws.cell(row=1, column=2).fill = mk_fill(C_HEADER_MAIN)

            for vi, partner in enumerate(all_partners):
                col = start_col + (vi * NUM_VENDOR_COLS)
                ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + 2)
                c = ws.cell(row=1, column=col, value=partner.name)
                style(c, bg=C_HEADER_MAIN, fnt=mk_font(bold=True, color=C_WHITE, size=12),
                      aln=mk_align(), brd=mk_border())
                for extra in range(1, 3):
                    ws.cell(row=1, column=col + extra).fill = mk_fill(C_HEADER_MAIN)
                    ws.cell(row=1, column=col + extra).border = mk_border()

            c = ws.cell(row=1, column=total_col, value="Grand Total")
            style(c, bg=C_TOTAL_COL, fnt=mk_font(bold=True, color=C_WHITE, size=12),
                  aln=mk_align(), brd=mk_border("F0C8A0"))

            # ─── Row 2: Sub Headers ──────────────────────────────────────────────
            for i, h in enumerate(["Product", "UoM"], start=1):
                c = ws.cell(row=2, column=i, value=h)
                style(c, bg=C_HEADER_SUB, fnt=mk_font(bold=True, color=C_DARK),
                      aln=mk_align(), brd=mk_border())

            for vi in range(len(all_partners)):
                col = start_col + (vi * NUM_VENDOR_COLS)
                for j, h in enumerate(["Price", "Qty", "Qty × Price"]):
                    bg = C_QTY_PRICE_SUB if h == "Qty × Price" else C_HEADER_SUB
                    c = ws.cell(row=2, column=col + j, value=h)
                    style(c, bg=bg, fnt=mk_font(bold=True, color=C_DARK),
                          aln=mk_align(), brd=mk_border())

            c = ws.cell(row=2, column=total_col, value="Grand Total")
            style(c, bg=C_TOTAL_SUB, fnt=mk_font(bold=True, color=C_DARK),
                  aln=mk_align(), brd=mk_border("F0C8A0"))

            # ─── Data Rows ───────────────────────────────────────────────────────
            for pi, (product, pdata) in enumerate(product_dict.items()):
                row = 3 + pi
                bg = C_ROW_ODD if pi % 2 == 0 else C_ROW_EVEN
                vendors = pdata['vendors']

                # Product name
                c = ws.cell(row=row, column=1, value=product.name)
                style(c, bg=bg, fnt=mk_font(bold=True, color=C_DARK),
                      aln=mk_align(h="left"), brd=mk_border())

                # UoM
                c = ws.cell(row=row, column=2, value=pdata['uom'])
                style(c, bg=bg, fnt=mk_font(color=C_GRAY),
                      aln=mk_align(), brd=mk_border())

                row_total_parts = []

                for vi, partner in enumerate(all_partners):
                    col = start_col + (vi * NUM_VENDOR_COLS)
                    vdata = vendors.get(partner)

                    if vdata:
                        # Price
                        pc = ws.cell(row=row, column=col, value=vdata['price'])
                        style(pc, bg=bg, fnt=mk_font(bold=True, color=C_PRICE),
                              aln=mk_align(), brd=mk_border())
                        pc.number_format = '#,##0.00'

                        # Qty
                        qc = ws.cell(row=row, column=col + 1, value=vdata['qty'])
                        style(qc, bg=bg, fnt=mk_font(bold=True, color=C_QTY),
                              aln=mk_align(), brd=mk_border())

                        # Qty × Price formula
                        pr = f"{get_column_letter(col)}{row}"
                        qr = f"{get_column_letter(col + 1)}{row}"
                        xc = ws.cell(row=row, column=col + 2, value=f"={pr}*{qr}")
                        style(xc, bg=bg, fnt=mk_font(bold=True, color=C_QXPRICE),
                              aln=mk_align(), brd=mk_border())
                        xc.number_format = '#,##0.00'
                        row_total_parts.append(f"{get_column_letter(col + 2)}{row}")
                    else:
                        for offset in range(3):
                            c = ws.cell(row=row, column=col + offset, value="-")
                            style(c, bg=bg, fnt=mk_font(color="BDC3C7"),
                                  aln=mk_align(), brd=mk_border())

                # Grand Total per row
                formula = ("=" + "+".join(row_total_parts)) if row_total_parts else 0
                c = ws.cell(row=row, column=total_col, value=formula)
                style(c, bg="FEF5EC", fnt=mk_font(bold=True, color=C_GRAND_TOTAL),
                      aln=mk_align(), brd=mk_border("F0C8A0"))
                c.number_format = '#,##0.00'

            # ─── Footer Total Row ────────────────────────────────────────────────
            footer_row = 3 + len(product_dict)

            ws.merge_cells(start_row=footer_row, start_column=1,
                           end_row=footer_row, end_column=2)
            c = ws.cell(row=footer_row, column=1, value="Total")
            style(c, bg=C_FOOTER_BG, fnt=mk_font(bold=True, color=C_DARK, size=12),
                  aln=mk_align(), brd=mk_border())
            ws.cell(row=footer_row, column=2).fill = mk_fill(C_FOOTER_BG)
            ws.cell(row=footer_row, column=2).border = mk_border()

            for vi in range(len(all_partners)):
                col = start_col + (vi * NUM_VENDOR_COLS)

                # Empty Price
                c = ws.cell(row=footer_row, column=col, value="")
                style(c, bg=C_FOOTER_BG, brd=mk_border())

                # Empty Qty
                c = ws.cell(row=footer_row, column=col + 1, value="")
                style(c, bg=C_FOOTER_BG, brd=mk_border())

                # SUM Qty×Price
                qxp_col = get_column_letter(col + 2)
                c = ws.cell(row=footer_row, column=col + 2,
                            value=f"=SUM({qxp_col}3:{qxp_col}{footer_row - 1})")
                style(c, bg=C_FOOTER_BG, fnt=mk_font(bold=True, color=C_QXPRICE, size=12),
                      aln=mk_align(), brd=mk_border())
                c.number_format = '#,##0.00'

            # Grand Total footer
            gt_col = get_column_letter(total_col)
            c = ws.cell(row=footer_row, column=total_col,
                        value=f"=SUM({gt_col}3:{gt_col}{footer_row - 1})")
            style(c, bg="FDE8D8", fnt=mk_font(bold=True, color=C_GRAND_TOTAL, size=12),
                  aln=mk_align(), brd=mk_border("F0C8A0"))
            c.number_format = '#,##0.00'

            # ─── Column Widths & Row Heights ─────────────────────────────────────
            ws.column_dimensions['A'].width = 24
            ws.column_dimensions['B'].width = 10

            for vi in range(len(all_partners)):
                col = start_col + (vi * NUM_VENDOR_COLS)
                ws.column_dimensions[get_column_letter(col)].width = 12
                ws.column_dimensions[get_column_letter(col + 1)].width = 10
                ws.column_dimensions[get_column_letter(col + 2)].width = 14

            ws.column_dimensions[get_column_letter(total_col)].width = 14

            ws.row_dimensions[1].height = 28
            ws.row_dimensions[2].height = 22
            for i in range(len(product_dict)):
                ws.row_dimensions[3 + i].height = 20
            ws.row_dimensions[footer_row].height = 24

            ws.freeze_panes = "A3"

            # ─── Save & Return Download ──────────────────────────────────────────
            buffer = io.BytesIO()
            wb.save(buffer)
            buffer.seek(0)
            file_data = base64.b64encode(buffer.read())

            attachment = self.env['ir.attachment'].create({
                'name': f'Vendor_Comparison_{rec.name}.xlsx',
                'type': 'binary',
                'datas': file_data,
                'mimetype': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            })

            return {
                'type': 'ir.actions.act_url',
                'url': f'/web/content/{attachment.id}?download=true',
                'target': 'self',
            }
    @api.depends('line_ids')
    def _compute_product_count(self):
        for rec in self:
            rec.product_count = len(rec.line_ids)

    @api.depends('line_ids.partner_ids')
    def _compute_vendor_count(self):
        for rec in self:
            all_vendors = rec.line_ids.mapped('partner_ids')
            rec.vendor_count = len(all_vendors)

    @api.depends('line_ids')
    def _compute_purchase_order(self):
        for rec in self:
            rec.purchase_count = self.env['purchase.order'].search_count([('vendor_comparison_id', '=', rec.id)])

    @api.model_create_multi
    def create(self, vals_list):
        for vals in vals_list:
            if vals.get('name', 'New') == 'New':
                vals['name'] = self.env['ir.sequence'].next_by_code('vendor.comparison.seq') or 'New'
        return super().create(vals_list)



    def action_create_rfq(self):
        self.ensure_one()
        if not self.line_ids:
            raise ValidationError("There are no product lines to create Purchase Orders.")
        for line in self.line_ids:
            if not line.product_id:
                raise ValidationError("Please select a Product in all lines.")
            if not line.partner_ids:
                raise ValidationError(
                    f"Please select at least one Vendor for product {line.product_id.display_name}."
                )
            if not line.product_qty or line.product_qty <= 0:
                raise ValidationError(
                    f"Quantity must be greater than 0 for product {line.product_id.display_name}."
                )
        vendor_map = {}
        for line in self.line_ids:  # Use all lines, no filtering
            for vendor in line.partner_ids:
                vendor_map.setdefault(vendor, []).append(line)
        created_pos = []
        for vendor, lines in vendor_map.items():
            new_po = self.env['purchase.order'].create({
                'partner_id': vendor.id,
                'origin': self.name,
                'vendor_comparison_id': self.id,
            })
            for line in lines:
                self.env['purchase.order.line'].create({
                    'order_id': new_po.id,
                    'product_id': line.product_id.id,
                    'product_qty': line.product_qty,
                    'product_uom_id': line.product_uom_id.id,
                    'price_unit': 0.0,
                })
            created_pos.append(new_po.id)
        self.state = 'rfq_created'
        return {
            'type': 'ir.actions.act_window',
            'res_model': 'purchase.order',
            'view_mode': 'list,form',
            'res_id': created_pos[0] if len(created_pos) == 1 else False,
            'domain': [('id', 'in', created_pos)] if len(created_pos) > 1 else [],
            'target': 'current',
        }
    def action_view_purchase(self):
        return {
            'type': 'ir.actions.act_window',
            'name': 'Tender RFQs',
            'res_model': 'purchase.order',
            'view_mode': 'list,form',
            'domain': [('vendor_comparison_id', '=', self.id)],
        }



class VendorComparisonLine(models.Model):
    _name = "vendor.comparison.line"
    _description = "Vendor Comparison Line"

    comparison_id = fields.Many2one('vendor.comparison', ondelete='cascade')
    partner_ids = fields.Many2many('res.partner', string="Vendors")
    product_id = fields.Many2one('product.product')
    name = fields.Text('Description')
    product_qty = fields.Float('Demand')
    product_uom_id = fields.Many2one('uom.uom', related='product_id.uom_id', string='UOM')
